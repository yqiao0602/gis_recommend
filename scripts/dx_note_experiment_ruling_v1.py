#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Append the 2026-10-17 experiment-design ruling (A1/A2/B single-variable
separation) to the workbook's methodology sheet. Idempotent: skips if the
marker row already exists."""
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

XLSX = Path(__file__).resolve().parents[1] / (
    "docs/reports/route2_p0_contract/plan_pipeline_v1/"
    "planning_capability_audit_v1/diag50_v1/dx_v1/dx50_v1/dx50_full_v1/"
    "DX50_S0_REVIEW_DX50_FULL_V1.xlsx")
MARKER = "实验设计细化裁定"


def main():
    wb = load_workbook(XLSX)
    wrap = Alignment(wrap_text=True, vertical="top")
    ws = wb["判读说明"]
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value and str(row[0].value).startswith(MARKER):
            print("already present, skip")
            return
    ws.append([
        MARKER,
        "owner 2026-10-17 晚间裁定（全文 DX_EXPERIMENT_DESIGN_RULING_OWNER_V1.md）："
        "S1 实为 ID 词面检索后，S2 弃权不可简单归因模型能力——三问分离 "
        "A1 卡片补全（冻结 S1 检索文本+Top-10，唯一变量=卡）/ A2 检索语料升级"
        "（唯一变量=ID 文本 vs layer_c_intent.text）/ B 规划能力对照（Base vs 训练版本），"
        "先分别做再测组合；区分「在库中/被 Top-10 召回/卡片足以证明适用」三态不合并。"
        "执行序：先做已放行 OTB 115 卡补全（A1 证据底座），补卡前哈希冻结已落"
        "（A_TRACK_P0_HASH_FREEZE_V1.json：profile v0_4 SHA 95ae97bc…，"
        "desc 键 0/1291，49 个 retrieval.json 全冻结）。"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).alignment = wrap
    try:
        wb.save(XLSX)
        print("appended:", XLSX)
    except PermissionError:
        print("workbook locked (open in Excel?) — not saved, rerun when closed")


if __name__ == "__main__":
    main()
